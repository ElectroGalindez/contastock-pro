"""
Tests del módulo de inventario y exportaciones a Excel.
"""
import pytest
from backend.productos import (
    guardar_producto, get_product, update_product, adjust_stock,
    eliminar_producto, list_products, exportar_inventario_excel,
)
from backend.categorias import agregar_categoria, list_categories


def _cat_id():
    agregar_categoria("Cat")
    return list_categories()[-1]["id"]


def test_crear_producto():
    prod = guardar_producto("Test A", 50.0, 5, _cat_id(), "test")
    assert prod["nombre"] == "Test A"
    assert float(prod["precio"]) == 50.0
    assert prod["cantidad"] == 5


def test_editar_producto_mismo_nombre(cliente, producto):
    # guardar_producto actualiza si el nombre ya existe
    actualizado = guardar_producto("Producto Test", 999.0, 3, producto["categoria_id"], "test")
    assert float(actualizado["precio"]) == 999.0
    assert actualizado["cantidad"] == 3


def test_ajustar_stock():
    prod = guardar_producto("Stock", 10.0, 10, _cat_id(), "test")
    adjust_stock(prod["id"], 5, "test")
    assert get_product(prod["id"])["cantidad"] == 15
    adjust_stock(prod["id"], -3, "test")
    assert get_product(prod["id"])["cantidad"] == 12


def test_eliminar_producto(cliente, producto):
    pid = producto["id"]
    eliminar_producto(pid, "test")
    assert get_product(pid) is None


def test_exportar_inventario_excel_valido(cliente, producto):
    data = exportar_inventario_excel()
    assert data[:2] == b"PK"  # firma de archivo xlsx

    from openpyxl import load_workbook
    from io import BytesIO
    ws = load_workbook(BytesIO(data)).active
    headers = [c.value for c in ws[1]]
    assert headers == ["Nombre", "Cantidad"]

    filas_datos = ws.max_row - 1
    assert filas_datos >= 1

    # El producto de prueba debe estar
    nombres = {ws.cell(row=r, column=1).value for r in range(2, ws.max_row + 1)}
    assert "Producto Test" in nombres


def test_list_products_retorna_todos(cliente, producto):
    guardar_producto("Extra", 10.0, 1, producto["categoria_id"], "test")
    prods = list_products()
    nombres = {p["nombre"] for p in prods}
    assert "Producto Test" in nombres
    assert "Extra" in nombres
"""
Tests del reporte "Ventas del Dia": el rango de fechas (fecha_inicio / fecha_fin)
debe filtrar las ventas correctamente y, por defecto, mostrar solo el dia de hoy.
Se usa un cliente distinto por venta para poder identificar cada una en el HTML/Excel.
"""
import io
from datetime import date, timedelta

import pytest
from openpyxl import load_workbook
from sqlalchemy import text

from app import create_app
from backend import db
from backend.clientes import add_client
from backend.ventas import register_sale, exportar_ventas_dia_excel


@pytest.fixture
def cliente_sesion():
    """Cliente de Flask autenticado como admin."""
    app = create_app()
    app.config.update(TESTING=True)
    with app.test_client() as c:
        with c.session_transaction() as s:
            s["usuario"] = {"id": 1, "username": "test", "rol": "admin"}
        yield c


def _venta(producto, dias_atras, nombre, total=100, pagado=100):
    """Registra una venta de un cliente propio y la fecha a `dias_atras` dias atras."""
    cli = add_client(nombre=nombre, telefono="5550000", ci=None, direccion=None, chapa=None)
    venta = register_sale(
        cliente_id=cli["id"], total=total, pagado=pagado,
        usuario="test", tipo_pago="Efectivo",
        productos=[{"id_producto": producto["id"], "nombre": producto["nombre"],
                    "cantidad": 1, "precio_unitario": total}],
    )
    fecha = date.today() - timedelta(days=dias_atras)
    with db.engine.begin() as conn:
        conn.execute(
            text("UPDATE ventas SET fecha = :fecha WHERE id = :id"),
            {"fecha": f"{fecha.isoformat()} 12:00:00", "id": venta["id"]},
        )
    return venta


def test_por_defecto_muestra_solo_hoy(cliente_sesion, producto):
    """Sin parametros, el reporte solo muestra las ventas del dia de hoy."""
    _venta(producto, 0, "Cliente Hoy")
    _venta(producto, 1, "Cliente Ayer")

    r = cliente_sesion.get("/ventas_dia")
    assert r.status_code == 200
    cuerpo = r.get_data(as_text=True)
    assert f'value="{date.today().isoformat()}"' in cuerpo
    assert "Cliente Hoy" in cuerpo
    assert "Cliente Ayer" not in cuerpo


def test_filtra_por_rango_incluido(cliente_sesion, producto):
    """Un rango de N dias debe incluir exactamente las ventas de esa ventana."""
    _venta(producto, 1, "Cliente Ayer")
    _venta(producto, 3, "Cliente Hace3")
    _venta(producto, 10, "Cliente Hace10")

    hoy = date.today()
    r = cliente_sesion.get("/ventas_dia", query_string={
        "fecha_inicio": (hoy - timedelta(days=4)).isoformat(),
        "fecha_fin": hoy.isoformat(),
    })
    assert r.status_code == 200
    cuerpo = r.get_data(as_text=True)
    assert "Cliente Ayer" in cuerpo
    assert "Cliente Hace3" in cuerpo
    assert "Cliente Hace10" not in cuerpo


def test_rango_de_un_solo_dia(cliente_sesion, producto):
    """Un rango con inicio == fin muestra solo ese dia."""
    _venta(producto, 5, "Cliente Dia5")
    _venta(producto, 6, "Cliente Dia6")

    dia = (date.today() - timedelta(days=5)).isoformat()
    r = cliente_sesion.get("/ventas_dia", query_string={"fecha_inicio": dia, "fecha_fin": dia})
    cuerpo = r.get_data(as_text=True)
    assert "Cliente Dia5" in cuerpo
    assert "Cliente Dia6" not in cuerpo


def test_rango_invertido_no_revienta(cliente_sesion, producto):
    """Si fecha_inicio > fecha_fin, avisa y recorta el rango a la fecha fin."""
    _venta(producto, 0, "Cliente Hoy")
    hoy = date.today()
    manana = (hoy + timedelta(days=1)).isoformat()

    r = cliente_sesion.get("/ventas_dia",
                           query_string={"fecha_inicio": manana, "fecha_fin": hoy.isoformat()})
    assert r.status_code == 200
    cuerpo = r.get_data(as_text=True)
    assert "no puede ser mayor" in cuerpo
    assert f'value="{hoy.isoformat()}"' in cuerpo
    assert "Cliente Hoy" in cuerpo


def test_fechas_invalidas_caen_en_hoy(cliente_sesion, producto):
    """Fechas no parseables vuelven al comportamiento por defecto (hoy)."""
    _venta(producto, 0, "Cliente Hoy")

    r = cliente_sesion.get("/ventas_dia", query_string={"fecha_inicio": "ayer", "fecha_fin": ""})
    assert r.status_code == 200
    cuerpo = r.get_data(as_text=True)
    assert f'value="{date.today().isoformat()}"' in cuerpo
    assert "Cliente Hoy" in cuerpo


def test_formulario_expone_los_dos_inputs(cliente_sesion, producto):
    """El formulario de filtro debe tener los inputs que usa descargarExcel()."""
    r = cliente_sesion.get("/ventas_dia")
    cuerpo = r.get_data(as_text=True)
    assert 'name="fecha_inicio"' in cuerpo
    assert 'name="fecha_fin"' in cuerpo


def test_excel_respeta_el_rango(producto):
    """El Excel generado usa el mismo rango inclusivo que la pantalla."""
    _venta(producto, 2, "Cliente Dentro")
    _venta(producto, 8, "Cliente Fuera")

    hoy = date.today()
    datos = exportar_ventas_dia_excel(
        (hoy - timedelta(days=3)).isoformat(), hoy.isoformat()
    )
    ws = load_workbook(io.BytesIO(datos)).active
    cuerpo = "\n".join(
        " ".join(str(c.value) for c in fila if c.value is not None)
        for fila in ws.iter_rows()
    )
    assert "Cliente Dentro" in cuerpo
    assert "Cliente Fuera" not in cuerpo


def test_excel_no_creasha_con_fecha_invalida():
    """Una fecha corrupta no debe levantar ValueError al exportar."""
    assert exportar_ventas_dia_excel("no-es-fecha", "2020-13-45")


# ----------------------------
# Orden del Detalle de Ventas
# ----------------------------
def test_detalle_ventas_ordena_por_id_descendente(cliente_sesion, producto):
    """El Detalle de Ventas va de mayor a menor id de venta, como /ventas."""
    # ids 1..4 con fechas en desorden, para distinguir orden por id de por fecha
    v1 = _venta(producto, 3, "Cliente Mas Viejo")    # id 1
    v2 = _venta(producto, 0, "Cliente Hoy")          # id 2, el mas reciente por fecha
    v3 = _venta(producto, 1, "Cliente Ayer")         # id 3
    v4 = _venta(producto, 2, "Cliente Hace2")        # id 4, el mayor id

    hoy = date.today()
    r = cliente_sesion.get("/ventas_dia", query_string={
        "fecha_inicio": (hoy - timedelta(days=5)).isoformat(),
        "fecha_fin": hoy.isoformat(),
    })
    assert r.status_code == 200
    cuerpo = r.get_data(as_text=True)

    # los cuatro tienen que estar en la pagina
    for c in ("Cliente Mas Viejo", "Cliente Hoy", "Cliente Ayer", "Cliente Hace2"):
        assert c in cuerpo, f"falta {c}"

    # el de mayor id va primero
    assert cuerpo.index("Cliente Hace2") < cuerpo.index("Cliente Ayer")
    assert cuerpo.index("Cliente Ayer") < cuerpo.index("Cliente Hoy")
    assert cuerpo.index("Cliente Hoy") < cuerpo.index("Cliente Mas Viejo")
    # no es orden por fecha: la mas reciente por fecha (Hoy) NO va primera
    assert cuerpo.index("Cliente Hoy") > cuerpo.index("Cliente Hace2")
    assert v1["id"] < v2["id"] < v3["id"] < v4["id"]


def test_excel_detalle_ventas_also_por_id_descendente(producto):
    """El Excel usa el mismo orden que la pantalla."""
    _venta(producto, 3, "Cliente Mas Viejo")
    _venta(producto, 0, "Cliente Hoy")
    _venta(producto, 1, "Cliente Ayer")
    _venta(producto, 2, "Cliente Hace2")

    hoy = date.today().isoformat()
    inicio = (date.today() - timedelta(days=5)).isoformat()
    data = exportar_ventas_dia_excel(inicio, hoy)
    wb = load_workbook(io.BytesIO(data))
    ws = wb.active

    # saltar la fila de encabezado, whose col 1 is exactly "Cliente"
    orden = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        if row and isinstance(row[1], str) and row[1].startswith("Cliente ") \
                and row[1] not in orden:
            orden.append(row[1])
    assert orden == ["Cliente Hace2", "Cliente Ayer", "Cliente Hoy", "Cliente Mas Viejo"]

    # OJO: el Excel NO tiene columna de id (solo Fecha/Cliente/Producto/
    # Cantidad/Estado), asi que el orden se verifica por el nombre del
    # cliente, que es una fila por venta. No tiene sentido buscar ints.
    assert [c.value for c in ws[1]][:5] == ["Fecha", "Cliente", "Producto", "Cantidad", "Estado"]
